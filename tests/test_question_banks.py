import json
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import Workbook, load_workbook
from PIL import Image

from quiz.question_banks import (
    QuestionBankImportConflict,
    QuestionBankValidationError,
    convert_legacy_question_bank,
    export_question_audit_batches,
    import_question_bank,
    validate_question_bank,
)

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]


def _write_minimal_bank(package_dir: Path) -> Path:
    package_dir.mkdir()
    (package_dir / "assets").mkdir()

    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "metadata"
    metadata.append(("key", "value"))
    metadata.append(("format_version", "1"))
    metadata.append(("bank_version", "example-v1"))
    metadata.append(("title", "虚构示例题库"))
    metadata.append(("source_label", "测试生成内容"))

    questions = workbook.create_sheet("questions")
    questions.append(
        (
            "stable_id",
            "category",
            "sampling_pool",
            "question_type",
            "prompt",
            "option_a",
            "option_b",
            "option_c",
            "option_d",
            "correct_option",
            "accepted_answers",
            "asset_path",
            "origin",
            "prior_event_use",
            "review_status",
            "private_event_allowed",
            "public_release_allowed",
            "source_reference",
            "source_license",
        )
    )
    questions.append(
        (
            "EXAMPLE-0001",
            "demo",
            "default",
            "multiple_choice",
            "虚构问题",
            "甲",
            "乙",
            "丙",
            "丁",
            "A",
            "",
            "",
            "synthetic_example",
            False,
            "verified",
            True,
            True,
            "generated-for-tests",
            "CC0-1.0",
        )
    )

    sampling_rules = workbook.create_sheet("sampling_rules")
    sampling_rules.append(("category", "sampling_pool", "default_quota", "display_order"))
    sampling_rules.append(("demo", "default", 1, 1))
    workbook.save(package_dir / "workbook.xlsx")
    return package_dir


def test_validate_question_bank_accepts_minimal_three_sheet_package(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")

    validated = validate_question_bank(package)

    assert validated.version == "example-v1"
    assert validated.content_hash.startswith("sha256:")
    assert len(validated.questions) == 1
    assert validated.questions[0].stable_id == "EXAMPLE-0001"


def test_public_example_is_a_valid_synthetic_package() -> None:
    example = validate_question_bank(REPOSITORY_ROOT / "content" / "examples" / "question-bank-v1")

    assert example.version == "public-example-v1"
    assert {question.origin for question in example.questions} == {"synthetic_example"}
    assert all(question.private_event_allowed for question in example.questions)
    assert all(question.public_release_allowed for question in example.questions)


def test_legacy_conversion_command_refuses_non_private_output(tmp_path: Path) -> None:
    with pytest.raises(CommandError, match="must be below an _private directory"):
        call_command(
            "convert_legacy_question_bank",
            str(tmp_path / "legacy.xlsx"),
            str(tmp_path / "public-output"),
        )


def test_validate_question_bank_rejects_asset_path_escape(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    workbook = load_workbook(package / "workbook.xlsx")
    workbook["questions"]["L2"] = "../outside.png"
    workbook.save(package / "workbook.xlsx")

    with pytest.raises(QuestionBankValidationError) as error:
        validate_question_bank(package)

    assert "asset_path must stay below assets/" in str(error.value)


def test_validate_question_bank_rejects_zero_sampling_quota(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    workbook = load_workbook(package / "workbook.xlsx")
    workbook["sampling_rules"]["C2"] = 0
    workbook.save(package / "workbook.xlsx")

    with pytest.raises(QuestionBankValidationError) as error:
        validate_question_bank(package)

    assert "default_quota must be positive" in str(error.value)


@pytest.mark.parametrize(
    ("cell", "value", "expected_issue"),
    [
        ("J2", "Z", "correct_option must be A, B, C, or D"),
        ("R2", "", "source_reference is required"),
        ("S2", "", "source_license is required"),
        ("sampling_rules!C2", 2, "quota 2 exceeds 1 questions"),
    ],
)
def test_validate_question_bank_rejects_invalid_answers_provenance_and_quotas(
    tmp_path: Path, cell: str, value: object, expected_issue: str
) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    workbook = load_workbook(package / "workbook.xlsx")
    if "!" in cell:
        sheet_name, coordinate = cell.split("!", 1)
    else:
        sheet_name, coordinate = "questions", cell
    workbook[sheet_name][coordinate] = value
    workbook.save(package / "workbook.xlsx")

    with pytest.raises(QuestionBankValidationError) as error:
        validate_question_bank(package)

    assert expected_issue in str(error.value)


def test_validate_question_bank_rejects_duplicate_ids_and_corrupt_images(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    (package / "assets" / "broken.png").write_bytes(b"not an image")
    workbook = load_workbook(package / "workbook.xlsx")
    questions = workbook["questions"]
    questions["L2"] = "assets/broken.png"
    questions.append([cell.value for cell in questions[2]])
    workbook.save(package / "workbook.xlsx")

    with pytest.raises(QuestionBankValidationError) as error:
        validate_question_bank(package)

    assert "duplicate stable_id 'EXAMPLE-0001'" in str(error.value)
    assert "invalid image 'assets/broken.png'" in str(error.value)


def test_validate_question_bank_rejects_missing_image(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    workbook = load_workbook(package / "workbook.xlsx")
    workbook["questions"]["L2"] = "assets/missing.png"
    workbook.save(package / "workbook.xlsx")

    with pytest.raises(QuestionBankValidationError) as error:
        validate_question_bank(package)

    assert "missing asset 'assets/missing.png'" in str(error.value)


def test_multiple_questions_can_reuse_one_validated_asset(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    Image.new("RGB", (2, 2), "white").save(package / "assets" / "shared.png")
    workbook = load_workbook(package / "workbook.xlsx")
    questions = workbook["questions"]
    questions["L2"] = "assets/shared.png"
    second = [cell.value for cell in questions[2]]
    second[0] = "EXAMPLE-0002"
    second[4] = "另一道虚构问题"
    questions.append(second)
    workbook["sampling_rules"]["C2"] = 2
    workbook.save(package / "workbook.xlsx")

    validated = validate_question_bank(package)

    assert len(validated.questions) == 2


def test_convert_legacy_question_bank_marks_content_private_and_pending(tmp_path: Path) -> None:
    source = tmp_path / "legacy-2024.xlsx"
    source_assets = tmp_path / "legacy-assets"
    source_assets.mkdir()
    Image.new("RGB", (2, 2), "white").save(source_assets / "7.png")
    legacy = Workbook()
    sheet = legacy.active
    sheet.append(
        (
            "num",
            "module",
            "level",
            "form",
            "question",
            "graph",
            "A",
            "B",
            "C",
            "D",
            "answer",
            "graphnum",
        )
    )
    sheet.append(
        (1, "demo", "default", "multiple choice", "虚构旧题", 1, "甲", "乙", "丙", "丁", "B", 7)
    )
    legacy.save(source)

    report = convert_legacy_question_bank(
        source, tmp_path / "converted", assets_source_dir=source_assets
    )

    assert report.question_count == 1
    assert report.asset_count == 1
    converted = report.validation.questions[0]
    assert converted.stable_id == "LEGACY2024-0001"
    assert converted.origin == "legacy_2024"
    assert converted.prior_event_use is True
    assert converted.review_status == "pending"
    assert converted.private_event_allowed is True
    assert converted.public_release_allowed is False


def test_legacy_conversion_preserves_the_historical_history_sampling_quotas(
    tmp_path: Path,
) -> None:
    source = tmp_path / "legacy-2024.xlsx"
    legacy = Workbook()
    sheet = legacy.active
    sheet.append(
        (
            "num",
            "module",
            "level",
            "form",
            "question",
            "graph",
            "A",
            "B",
            "C",
            "D",
            "answer",
            "graphnum",
        )
    )
    for number in range(1, 30):
        sheet.append(
            (
                number,
                "history",
                "legacy-level",
                "multiple choice",
                f"虚构历史题 {number}",
                0,
                "甲",
                "乙",
                "丙",
                "丁",
                "A",
                "",
            )
        )
    legacy.save(source)

    converted = convert_legacy_question_bank(source, tmp_path / "converted")

    quotas = {
        rule.sampling_pool: rule.default_quota for rule in converted.validation.sampling_rules
    }
    assert quotas == {"legacy-1": 9, "legacy-2": 5, "legacy-3": 1}


def test_export_question_audit_batches_includes_review_placeholders(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")

    batches = export_question_audit_batches(package, tmp_path / "audit")

    assert len(batches) == 1
    records = [json.loads(line) for line in batches[0].read_text(encoding="utf-8").splitlines()]
    assert records[0]["stable_id"] == "EXAMPLE-0001"
    assert records[0]["review_request"] == {
        "evidence_urls": [],
        "verdict": "",
        "suggested_correction": "",
        "confidence": "",
    }
    index = json.loads((tmp_path / "audit" / "index.json").read_text(encoding="utf-8"))
    assert index["human_confirmation_required"] is True


@pytest.mark.django_db
def test_import_question_bank_is_idempotent_for_same_version_and_hash(tmp_path: Path) -> None:
    package = _write_minimal_bank(tmp_path / "bank")

    first = import_question_bank(package)
    repeated = import_question_bank(package)

    assert first.created is True
    assert repeated.created is False
    assert repeated.version.pk == first.version.pk
    assert first.version.is_finalized is True
    assert first.version.questions.get().identity.stable_code == "EXAMPLE-0001"
    assert first.version.questions.get().source_license == "CC0-1.0"


@pytest.mark.django_db
def test_import_question_bank_copies_validated_assets_to_managed_media(
    tmp_path: Path, settings
) -> None:
    package = _write_minimal_bank(tmp_path / "bank")
    Image.new("RGB", (2, 2), "white").save(package / "assets" / "diagram.png")
    workbook = load_workbook(package / "workbook.xlsx")
    workbook["questions"]["L2"] = "assets/diagram.png"
    workbook.save(package / "workbook.xlsx")
    settings.MEDIA_ROOT = tmp_path / "media"

    result = import_question_bank(package)

    assert result.created is True
    assert (
        settings.MEDIA_ROOT / "question-banks" / "example-v1" / "assets" / "diagram.png"
    ).is_file()


@pytest.mark.django_db
def test_import_question_bank_rejects_reused_version_with_different_content(tmp_path: Path) -> None:
    first_package = _write_minimal_bank(tmp_path / "first")
    changed_package = _write_minimal_bank(tmp_path / "changed")
    workbook = load_workbook(changed_package / "workbook.xlsx")
    workbook["questions"]["E2"] = "另一条虚构问题"
    workbook.save(changed_package / "workbook.xlsx")
    import_question_bank(first_package)

    with pytest.raises(QuestionBankImportConflict):
        import_question_bank(changed_package)


@pytest.mark.django_db
def test_import_question_bank_does_not_write_any_rows_when_validation_fails(tmp_path: Path) -> None:
    from quiz.models import QuestionBankVersion, QuestionIdentity

    package = _write_minimal_bank(tmp_path / "bank")
    workbook = load_workbook(package / "workbook.xlsx")
    workbook["questions"]["R2"] = ""
    workbook.save(package / "workbook.xlsx")

    with pytest.raises(QuestionBankValidationError):
        import_question_bank(package)

    assert QuestionBankVersion.objects.count() == 0
    assert QuestionIdentity.objects.count() == 0


@pytest.mark.django_db
def test_create_activity_edition_command_builds_categories_and_opens_activity(
    tmp_path: Path, settings
) -> None:
    from quiz.models import ActivityEdition, ActivityStatus, AdminAuditLog

    settings.MEDIA_ROOT = tmp_path / "media"
    example = REPOSITORY_ROOT / "content" / "examples" / "question-bank-v1"
    imported = import_question_bank(example)
    get_user_model().objects.create_user(username="operator")

    call_command(
        "create_activity_edition",
        "demo-2026",
        "虚构示例活动",
        imported.version.version_code,
        actor="operator",
        reason="测试初始化",
        open=True,
    )

    activity = ActivityEdition.objects.get(slug="demo-2026")
    assert activity.status == ActivityStatus.OPEN
    assert activity.is_participant_entry is True
    assert list(activity.category_configs.values_list("category_key", flat=True)) == ["demo"]
    assert activity.bank_activations.get(is_current=True).bank == imported.version
    assert set(AdminAuditLog.objects.values_list("action", flat=True)) == {
        "question_bank_activated",
        "participant_entry_selected",
        "activity_status_changed",
    }
