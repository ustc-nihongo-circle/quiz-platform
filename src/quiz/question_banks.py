"""Validate, import, convert, and review versioned question-bank packages.

The public repository may contain only synthetic examples.  Callers are responsible for
keeping active and legacy packages under ``_private/``; the management commands enforce
that boundary for operations which copy real question content.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import Workbook, load_workbook
from PIL import Image, UnidentifiedImageError

FORMAT_VERSION = "1"
REQUIRED_SHEETS = ("metadata", "questions", "sampling_rules")
REVIEW_STATUSES = frozenset({"pending", "verified", "rejected"})
QUESTION_TYPES = frozenset({"multiple_choice", "blank_filling"})

QUESTION_COLUMNS = (
    "stable_id",
    "category",
    "sampling_pool",
    "question_type",
    "prompt",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
    "accepted_answers",
    "asset_path",
    "origin",
    "prior_event_use",
    "review_status",
    "private_event_allowed",
    "public_release_allowed",
    "source_reference",
    "source_license",
)
SAMPLING_RULE_COLUMNS = ("category", "sampling_pool", "default_quota", "display_order")


class QuestionBankValidationError(ValueError):
    """Raised with every detected package issue, before any database write occurs."""

    def __init__(self, issues: list[str] | tuple[str, ...]) -> None:
        self.issues = tuple(issues)
        super().__init__("Question bank validation failed:\n- " + "\n- ".join(self.issues))


class QuestionBankImportConflict(ValueError):
    """Raised when a version code is reused for different package content."""


@dataclass(frozen=True, slots=True)
class QuestionRecord:
    stable_id: str
    category: str
    sampling_pool: str
    question_type: str
    prompt: str
    options: tuple[str, str, str, str]
    correct_option: str
    accepted_answers: tuple[str, ...]
    asset_path: str
    origin: str
    prior_event_use: bool
    review_status: str
    private_event_allowed: bool
    public_release_allowed: bool
    source_reference: str
    source_license: str


@dataclass(frozen=True, slots=True)
class SamplingRuleRecord:
    category: str
    sampling_pool: str
    default_quota: int
    display_order: int


@dataclass(frozen=True, slots=True)
class ValidatedQuestionBank:
    package_dir: Path
    version: str
    title: str
    source_label: str
    content_hash: str
    metadata: dict[str, str]
    questions: tuple[QuestionRecord, ...]
    sampling_rules: tuple[SamplingRuleRecord, ...]


@dataclass(frozen=True, slots=True)
class ImportResult:
    version: Any
    created: bool


@dataclass(frozen=True, slots=True)
class LegacyConversionReport:
    package_dir: Path
    question_count: int
    asset_count: int
    warnings: tuple[str, ...]
    validation: ValidatedQuestionBank


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _boolean(value: Any, *, field: str, row_number: int, issues: list[str]) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _text(value).casefold()
    if normalized in {"true", "1", "yes"}:
        return True
    if normalized in {"false", "0", "no", ""}:
        return False
    issues.append(f"questions row {row_number}: {field} must be a boolean")
    return False


def _integer(value: Any, *, field: str, row_number: int, issues: list[str]) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError):
        issues.append(f"sampling_rules row {row_number}: {field} must be an integer")
        return 0
    if result < 0:
        issues.append(f"sampling_rules row {row_number}: {field} cannot be negative")
    return result


def _read_table(
    sheet: Any, expected_columns: tuple[str, ...], issues: list[str]
) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_header = next(rows)
    except StopIteration:
        issues.append(f"sheet {sheet.title!r} is empty")
        return []
    header = tuple(_text(value) for value in raw_header)
    if header != expected_columns:
        issues.append(
            f"sheet {sheet.title!r} columns must be exactly: {', '.join(expected_columns)}"
        )
        return []
    records = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        padded = tuple(values) + (None,) * (len(expected_columns) - len(values))
        records.append(dict(zip(expected_columns, padded, strict=False), _row_number=row_number))
    return records


def _read_metadata(sheet: Any, issues: list[str]) -> dict[str, str]:
    rows = sheet.iter_rows(values_only=True)
    try:
        header = tuple(_text(value) for value in next(rows))
    except StopIteration:
        issues.append("sheet 'metadata' is empty")
        return {}
    if header != ("key", "value"):
        issues.append("sheet 'metadata' columns must be exactly: key, value")
        return {}
    metadata: dict[str, str] = {}
    for row_number, values in enumerate(rows, start=2):
        key = _text(values[0] if values else None)
        value = _text(values[1] if len(values) > 1 else None)
        if not key and not value:
            continue
        if not key:
            issues.append(f"metadata row {row_number}: key is required")
        elif key in metadata:
            issues.append(f"metadata row {row_number}: duplicate key {key!r}")
        else:
            metadata[key] = value
    for key in ("format_version", "bank_version", "title", "source_label"):
        if not metadata.get(key):
            issues.append(f"metadata: {key} is required")
    if metadata.get("format_version") not in (None, "", FORMAT_VERSION):
        issues.append(
            f"metadata: unsupported format_version {metadata['format_version']!r}; "
            f"expected {FORMAT_VERSION!r}"
        )
    return metadata


def _safe_asset(
    package_dir: Path, asset_path: str, row_number: int, issues: list[str]
) -> Path | None:
    logical_path = PurePosixPath(asset_path)
    if (
        logical_path.is_absolute()
        or ".." in logical_path.parts
        or not logical_path.parts
        or logical_path.parts[0] != "assets"
        or "\\" in asset_path
    ):
        issues.append(
            f"questions row {row_number}: asset_path must stay below assets/ using '/' separators"
        )
        return None
    package_root = package_dir.resolve()
    candidate = (package_root / Path(*logical_path.parts)).resolve()
    try:
        candidate.relative_to((package_root / "assets").resolve())
    except ValueError:
        issues.append(f"questions row {row_number}: asset_path escapes the package")
        return None
    if not candidate.is_file():
        issues.append(f"questions row {row_number}: missing asset {asset_path!r}")
        return None
    try:
        with Image.open(candidate) as image:
            image.verify()
    except (OSError, UnidentifiedImageError):
        issues.append(f"questions row {row_number}: invalid image {asset_path!r}")
        return None
    return candidate


def _package_hash(package_dir: Path) -> str:
    digest = hashlib.sha256()
    files = [package_dir / "workbook.xlsx"]
    assets_dir = package_dir / "assets"
    if assets_dir.is_dir():
        files.extend(path for path in assets_dir.rglob("*") if path.is_file())
    for path in sorted(files, key=lambda item: item.relative_to(package_dir).as_posix()):
        relative = path.relative_to(package_dir).as_posix().encode("utf-8")
        digest.update(len(relative).to_bytes(4, "big"))
        digest.update(relative)
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def validate_question_bank(package_dir: str | Path) -> ValidatedQuestionBank:
    """Read and fully validate a three-sheet question-bank directory package."""

    package = Path(package_dir)
    issues: list[str] = []
    workbook_path = package / "workbook.xlsx"
    assets_dir = package / "assets"
    if not package.is_dir():
        raise QuestionBankValidationError([f"package directory does not exist: {package}"])
    if not workbook_path.is_file():
        raise QuestionBankValidationError(["package must contain workbook.xlsx"])
    if not assets_dir.is_dir():
        issues.append("package must contain an assets/ directory")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exposes several malformed-archive error types
        raise QuestionBankValidationError([f"cannot read workbook.xlsx: {exc}"]) from exc

    try:
        if tuple(workbook.sheetnames) != REQUIRED_SHEETS:
            issues.append("workbook sheets must be exactly: metadata, questions, sampling_rules")
        metadata = (
            _read_metadata(workbook["metadata"], issues)
            if "metadata" in workbook.sheetnames
            else {}
        )
        raw_questions = (
            _read_table(workbook["questions"], QUESTION_COLUMNS, issues)
            if "questions" in workbook.sheetnames
            else []
        )
        raw_rules = (
            _read_table(workbook["sampling_rules"], SAMPLING_RULE_COLUMNS, issues)
            if "sampling_rules" in workbook.sheetnames
            else []
        )
    finally:
        workbook.close()

    if not raw_questions:
        issues.append("questions sheet must contain at least one question")
    if not raw_rules:
        issues.append("sampling_rules sheet must contain at least one rule")

    questions: list[QuestionRecord] = []
    stable_ids: set[str] = set()
    pool_counts: dict[tuple[str, str], int] = {}
    for raw in raw_questions:
        row_number = raw.pop("_row_number")
        values = {key: _text(value) for key, value in raw.items()}
        stable_id = values["stable_id"]
        category = values["category"]
        sampling_pool = values["sampling_pool"]
        question_type = values["question_type"]
        prompt = values["prompt"]
        for field, value in (
            ("stable_id", stable_id),
            ("category", category),
            ("sampling_pool", sampling_pool),
            ("question_type", question_type),
            ("prompt", prompt),
            ("origin", values["origin"]),
            ("review_status", values["review_status"]),
            ("source_reference", values["source_reference"]),
            ("source_license", values["source_license"]),
        ):
            if not value:
                issues.append(f"questions row {row_number}: {field} is required")
        if stable_id in stable_ids:
            issues.append(f"questions row {row_number}: duplicate stable_id {stable_id!r}")
        stable_ids.add(stable_id)
        if question_type not in QUESTION_TYPES:
            issues.append(
                f"questions row {row_number}: question_type must be "
                "multiple_choice or blank_filling"
            )
        options = tuple(values[f"option_{letter}"] for letter in "abcd")
        correct_option = values["correct_option"].upper()
        accepted_answers = tuple(
            answer.strip() for answer in values["accepted_answers"].splitlines() if answer.strip()
        )
        if question_type == "multiple_choice":
            if not all(options):
                issues.append(f"questions row {row_number}: multiple choice requires options A-D")
            if correct_option not in {"A", "B", "C", "D"}:
                issues.append(
                    f"questions row {row_number}: correct_option must be A, B, C, or D"
                )
            if accepted_answers:
                issues.append(
                    f"questions row {row_number}: multiple choice cannot have accepted_answers"
                )
        elif question_type == "blank_filling":
            if any(options) or correct_option:
                issues.append(
                    f"questions row {row_number}: blank filling cannot have options "
                    "or correct_option"
                )
            if not accepted_answers:
                issues.append(
                    f"questions row {row_number}: blank filling requires accepted_answers"
                )

        review_status = values["review_status"]
        if review_status not in REVIEW_STATUSES:
            issues.append(f"questions row {row_number}: invalid review_status {review_status!r}")
        prior_event_use = _boolean(
            raw["prior_event_use"], field="prior_event_use", row_number=row_number, issues=issues
        )
        private_event_allowed = _boolean(
            raw["private_event_allowed"],
            field="private_event_allowed",
            row_number=row_number,
            issues=issues,
        )
        public_release_allowed = _boolean(
            raw["public_release_allowed"],
            field="public_release_allowed",
            row_number=row_number,
            issues=issues,
        )
        if public_release_allowed and review_status != "verified":
            issues.append(
                f"questions row {row_number}: public_release_allowed requires verified review"
            )

        asset_path = values["asset_path"]
        if asset_path:
            _safe_asset(package, asset_path, row_number, issues)
        pool_key = (category, sampling_pool)
        pool_counts[pool_key] = pool_counts.get(pool_key, 0) + 1
        questions.append(
            QuestionRecord(
                stable_id=stable_id,
                category=category,
                sampling_pool=sampling_pool,
                question_type=question_type,
                prompt=prompt,
                options=options,  # type: ignore[arg-type]
                correct_option=correct_option,
                accepted_answers=accepted_answers,
                asset_path=asset_path,
                origin=values["origin"],
                prior_event_use=prior_event_use,
                review_status=review_status,
                private_event_allowed=private_event_allowed,
                public_release_allowed=public_release_allowed,
                source_reference=values["source_reference"],
                source_license=values["source_license"],
            )
        )

    rules: list[SamplingRuleRecord] = []
    rule_keys: set[tuple[str, str]] = set()
    display_orders: set[tuple[str, int]] = set()
    for raw in raw_rules:
        row_number = raw.pop("_row_number")
        category = _text(raw["category"])
        sampling_pool = _text(raw["sampling_pool"])
        if not category:
            issues.append(f"sampling_rules row {row_number}: category is required")
        if not sampling_pool:
            issues.append(f"sampling_rules row {row_number}: sampling_pool is required")
        quota = _integer(
            raw["default_quota"], field="default_quota", row_number=row_number, issues=issues
        )
        if quota == 0:
            issues.append(f"sampling_rules row {row_number}: default_quota must be positive")
        display_order = _integer(
            raw["display_order"], field="display_order", row_number=row_number, issues=issues
        )
        key = (category, sampling_pool)
        if key in rule_keys:
            issues.append(f"sampling_rules row {row_number}: duplicate category and sampling_pool")
        rule_keys.add(key)
        display_key = (category, display_order)
        if display_key in display_orders:
            issues.append(
                f"sampling_rules row {row_number}: duplicate display_order within category"
            )
        display_orders.add(display_key)
        available = pool_counts.get(key, 0)
        if not available:
            issues.append(f"sampling_rules row {row_number}: pool has no questions")
        elif quota > available:
            issues.append(
                f"sampling_rules row {row_number}: quota {quota} exceeds {available} questions"
            )
        rules.append(SamplingRuleRecord(category, sampling_pool, quota, display_order))

    for key in sorted(pool_counts.keys() - rule_keys):
        issues.append(f"questions pool {key[0]!r}/{key[1]!r} has no sampling rule")
    if issues:
        raise QuestionBankValidationError(issues)

    return ValidatedQuestionBank(
        package_dir=package.resolve(),
        version=metadata["bank_version"],
        title=metadata["title"],
        source_label=metadata["source_label"],
        content_hash=_package_hash(package),
        metadata=metadata,
        questions=tuple(questions),
        sampling_rules=tuple(rules),
    )


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_question_bank(package_dir: str | Path) -> ImportResult:
    """Atomically import and finalize a validated package.

    Re-importing byte-identical content with the same version code is a no-op. Reusing the
    version code for different bytes is an explicit conflict, never an in-place update.
    """

    from django.conf import settings
    from django.db import transaction
    from django.utils import timezone

    from quiz.models import (
        BankQuestion,
        QuestionAsset,
        QuestionBankVersion,
        QuestionIdentity,
        QuestionType,
        SamplingRule,
    )

    bank = validate_question_bank(package_dir)
    digest = bank.content_hash.removeprefix("sha256:")
    with transaction.atomic():
        existing = (
            QuestionBankVersion.objects.select_for_update()
            .filter(version_code=bank.version)
            .first()
        )
        if existing is not None:
            if existing.content_sha256 == digest:
                return ImportResult(existing, False)
            raise QuestionBankImportConflict(
                f"question-bank version {bank.version!r} already exists with different content"
            )

        version = QuestionBankVersion.objects.create(
            version_code=bank.version,
            title=bank.title,
            format_version=int(bank.metadata["format_version"]),
            source_label=bank.source_label,
            content_sha256=digest,
        )
        imported_questions: dict[str, Any] = {}
        for question in bank.questions:
            identity, _ = QuestionIdentity.objects.get_or_create(stable_code=question.stable_id)
            values = {
                "bank": version,
                "identity": identity,
                "category_key": question.category,
                "pool_key": question.sampling_pool,
                "question_type": (
                    QuestionType.SINGLE_CHOICE
                    if question.question_type == "multiple_choice"
                    else QuestionType.FILL_BLANK
                ),
                "prompt": question.prompt,
                "option_a": question.options[0],
                "option_b": question.options[1],
                "option_c": question.options[2],
                "option_d": question.options[3],
                "correct_option": question.correct_option,
                "acceptable_answers": list(question.accepted_answers),
                "origin": question.origin,
                "prior_event_use": question.prior_event_use,
                "review_status": question.review_status,
                "private_event_allowed": question.private_event_allowed,
                "public_release_allowed": question.public_release_allowed,
                "source_reference": question.source_reference,
                "source_license": question.source_license,
            }
            imported_questions[question.stable_id] = BankQuestion.objects.create(**values)

        for question in bank.questions:
            if not question.asset_path:
                continue
            path = bank.package_dir / Path(*PurePosixPath(question.asset_path).parts)
            with Image.open(path) as image:
                mime_type = image.get_format_mimetype() or "application/octet-stream"
            QuestionAsset.objects.create(
                bank=version,
                question=imported_questions[question.stable_id],
                relative_path=question.asset_path,
                content_sha256=_file_sha256(path),
                mime_type=mime_type,
            )

        SamplingRule.objects.bulk_create(
            [
                SamplingRule(
                    bank=version,
                    category_key=rule.category,
                    pool_key=rule.sampling_pool,
                    default_quota=rule.default_quota,
                    display_order=rule.display_order,
                )
                for rule in bank.sampling_rules
            ]
        )
        version.finalized_at = timezone.now()
        version.save(update_fields=("finalized_at",))
        media_root = Path(settings.MEDIA_ROOT) / "question-banks" / bank.version
        if media_root.exists():
            raise QuestionBankImportConflict(
                f"managed media already exists for question-bank version {bank.version!r}"
            )
        try:
            for question in bank.questions:
                if not question.asset_path:
                    continue
                logical_path = PurePosixPath(question.asset_path)
                source = bank.package_dir / Path(*logical_path.parts)
                destination = media_root / Path(*logical_path.parts)
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, destination)
        except Exception:
            if media_root.exists():
                shutil.rmtree(media_root)
            raise
    return ImportResult(version, True)


def _legacy_answer_variants(value: Any) -> str:
    """Preserve legacy answers while converting their line breaks to the v1 wire format."""

    return "\n".join(part.strip() for part in _text(value).splitlines() if part.strip())


def _new_bank_workbook(
    *, version: str, title: str, source_label: str
) -> tuple[Workbook, Any, Any]:
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "metadata"
    metadata.append(("key", "value"))
    metadata.append(("format_version", FORMAT_VERSION))
    metadata.append(("bank_version", version))
    metadata.append(("title", title))
    metadata.append(("source_label", source_label))
    questions = workbook.create_sheet("questions")
    questions.append(QUESTION_COLUMNS)
    sampling_rules = workbook.create_sheet("sampling_rules")
    sampling_rules.append(SAMPLING_RULE_COLUMNS)
    return workbook, questions, sampling_rules


LEGACY_SAMPLING_GROUPS = {
    "history": ((2, 19, "legacy-1", 9), (20, 25, "legacy-2", 5), (26, 30, "legacy-3", 1)),
    "geography": (
        (31, 60, "legacy-1", 8),
        (61, 80, "legacy-2", 6),
        (81, 84, "legacy-3", 1),
    ),
    "yakyuu": ((85, 104, "legacy-1", 14), (105, 108, "legacy-2", 1)),
    "horse": ((109, 123, "legacy-1", 15),),
    "japanese": (
        (124, 138, "legacy-1", 1),
        (139, 153, "legacy-2", 8),
        (154, 167, "legacy-3", 5),
        (168, 174, "legacy-4", 1),
    ),
    "ACG": (
        (175, 270, "legacy-1", 3),
        (271, 300, "legacy-2", 2),
        (301, 313, "legacy-3", 4),
        (314, 323, "legacy-4", 5),
        (324, 334, "legacy-5", 1),
    ),
    "mahjong": (
        (335, 350, "legacy-1", 2),
        (351, 361, "legacy-2", 7),
        (362, 379, "legacy-3", 6),
    ),
}


def _legacy_sampling_group(
    *, row_number: int, category: str, original_level: str, warnings: list[str]
) -> tuple[str, int, int]:
    groups = LEGACY_SAMPLING_GROUPS.get(category, ())
    for display_order, (start, end, pool, quota) in enumerate(groups, start=1):
        if start <= row_number <= end:
            return pool, quota, display_order
    warnings.append(
        f"row {row_number}: no preserved legacy sampling group for category {category!r}"
    )
    return original_level or "legacy-unmapped", 1, len(groups) + 1


def convert_legacy_question_bank(
    source_workbook: str | Path,
    output_dir: str | Path,
    *,
    assets_source_dir: str | Path | None = None,
    version: str = "legacy-2024-converted",
) -> LegacyConversionReport:
    """Convert the 2024 workbook without modifying it or approving any of its content."""

    source = Path(source_workbook)
    output = Path(output_dir)
    if not source.is_file():
        raise FileNotFoundError(source)
    if (output / "workbook.xlsx").exists() or (output / "assets").exists():
        raise FileExistsError(f"conversion target is not empty: {output}")
    output.mkdir(parents=True, exist_ok=True)
    output_assets = output / "assets"
    output_assets.mkdir()
    source_assets = (
        Path(assets_source_dir)
        if assets_source_dir is not None
        else source.parent.parent / "legacy-reference" / "static" / "img"
    )

    legacy = load_workbook(source, read_only=True, data_only=True)
    warnings: list[str] = []
    asset_names: set[str] = set()
    try:
        if len(legacy.sheetnames) != 1:
            raise QuestionBankValidationError(
                ["legacy workbook must contain exactly one question sheet"]
            )
        sheet = legacy[legacy.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = tuple(_text(value) for value in next(rows))
        except StopIteration as exc:
            raise QuestionBankValidationError(["legacy workbook is empty"]) from exc
        expected = (
            "num",
            "module",
            "level",
            "form",
            "question",
            "graph",
            "A",
            "B",
            "C",
            "D",
            "answer",
            "graphnum",
        )
        if header != expected:
            raise QuestionBankValidationError(
                ["legacy workbook columns do not match the 2024 format"]
            )

        workbook, questions_sheet, rules_sheet = _new_bank_workbook(
            version=version,
            title="2024 往届题库转换副本",
            source_label="legacy-2024.xlsx; pending human review",
        )
        pool_rules: dict[tuple[str, str], tuple[int, int]] = {}
        question_count = 0
        for row_number, raw_values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in raw_values):
                continue
            padded = tuple(raw_values) + (None,) * (len(expected) - len(raw_values))
            raw = dict(zip(expected, padded, strict=False))
            question_count += 1
            category = _text(raw["module"])
            sampling_pool, quota, display_order = _legacy_sampling_group(
                row_number=row_number,
                category=category,
                original_level=_text(raw["level"]),
                warnings=warnings,
            )
            pool_key = (category, sampling_pool)
            pool_rules.setdefault(pool_key, (quota, display_order))
            if not _text(raw["num"]).isdigit():
                warnings.append(
                    f"row {row_number}: legacy num {_text(raw['num'])!r} is not numeric"
                )
            legacy_form = _text(raw["form"]).casefold()
            if legacy_form == "multiple choice":
                question_type = "multiple_choice"
                options = tuple(_text(raw[letter]) for letter in "ABCD")
                correct_option = _text(raw["answer"]).upper()
                accepted_answers = ""
            elif legacy_form == "blank filling":
                question_type = "blank_filling"
                options = ("", "", "", "")
                correct_option = ""
                accepted_answers = _legacy_answer_variants(raw["answer"])
            else:
                question_type = legacy_form.replace(" ", "_")
                options = tuple(_text(raw[letter]) for letter in "ABCD")
                correct_option = _text(raw["answer"]).upper()
                accepted_answers = ""
                warnings.append(f"row {row_number}: unknown form {legacy_form!r}")

            graph_number = _text(raw["graphnum"])
            asset_path = ""
            if graph_number:
                source_asset = source_assets / f"{graph_number}.png"
                if not source_asset.is_file():
                    warnings.append(f"row {row_number}: missing legacy image {graph_number}.png")
                    asset_path = f"assets/{graph_number}.png"
                else:
                    asset_name = f"{graph_number}.png"
                    if asset_name not in asset_names:
                        shutil.copyfile(source_asset, output_assets / asset_name)
                        asset_names.add(asset_name)
                    asset_path = f"assets/{asset_name}"
            elif _text(raw["graph"]).casefold() in {"1", "true", "yes"}:
                warnings.append(f"row {row_number}: graph flag is set but graphnum is empty")

            questions_sheet.append(
                (
                    f"LEGACY2024-{question_count:04d}",
                    category,
                    sampling_pool,
                    question_type,
                    _text(raw["question"]),
                    *options,
                    correct_option,
                    accepted_answers,
                    asset_path,
                    "legacy_2024",
                    True,
                    "pending",
                    True,
                    False,
                    f"legacy-2024.xlsx row {row_number}",
                    "unreviewed",
                )
            )
    finally:
        legacy.close()

    for (category, sampling_pool), (quota, display_order) in pool_rules.items():
        rules_sheet.append(
            (category, sampling_pool, quota, display_order)
        )
    workbook.save(output / "workbook.xlsx")
    workbook.close()

    report_path = output / "conversion-report.json"
    report_path.write_text(
        json.dumps(
            {
                "source": source.name,
                "question_count": question_count,
                "asset_count": len(asset_names),
                "warnings": warnings,
                "review_status": "pending",
                "public_release_allowed": False,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    validation = validate_question_bank(output)
    return LegacyConversionReport(
        package_dir=output.resolve(),
        question_count=question_count,
        asset_count=len(asset_names),
        warnings=tuple(warnings),
        validation=validation,
    )


def _batch_filename(category: str, sampling_pool: str) -> str:
    label = f"{category}-{sampling_pool}"
    slug = re.sub(r"[^a-z0-9]+", "-", label.casefold()).strip("-") or "questions"
    suffix = hashlib.sha256(label.encode("utf-8")).hexdigest()[:8]
    return f"{slug}-{suffix}.jsonl"


def export_question_audit_batches(
    package_dir: str | Path, output_dir: str | Path
) -> tuple[Path, ...]:
    """Export private review batches with empty fields for model and human findings."""

    bank = validate_question_bank(package_dir)
    output = Path(output_dir)
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"audit target is not empty: {output}")
    output.mkdir(parents=True, exist_ok=True)
    grouped: dict[tuple[str, str], list[QuestionRecord]] = {}
    for question in bank.questions:
        grouped.setdefault((question.category, question.sampling_pool), []).append(question)

    paths: list[Path] = []
    index: list[dict[str, Any]] = []
    for (category, sampling_pool), questions in grouped.items():
        path = output / _batch_filename(category, sampling_pool)
        with path.open("w", encoding="utf-8", newline="\n") as destination:
            for question in questions:
                existing_answer: str | list[str]
                if question.question_type == "multiple_choice":
                    existing_answer = question.correct_option
                else:
                    existing_answer = list(question.accepted_answers)
                record = {
                    "stable_id": question.stable_id,
                    "category": category,
                    "sampling_pool": sampling_pool,
                    "question_type": question.question_type,
                    "prompt": question.prompt,
                    "options": list(question.options),
                    "existing_answer": existing_answer,
                    "asset_path": question.asset_path or None,
                    "source_reference": question.source_reference,
                    "review_status": question.review_status,
                    "review_request": {
                        "evidence_urls": [],
                        "verdict": "",
                        "suggested_correction": "",
                        "confidence": "",
                    },
                }
                destination.write(json.dumps(record, ensure_ascii=False) + "\n")
        paths.append(path)
        index.append(
            {
                "category": category,
                "sampling_pool": sampling_pool,
                "file": path.name,
                "question_count": len(questions),
            }
        )
    (output / "index.json").write_text(
        json.dumps(
            {
                "bank_version": bank.version,
                "content_hash": bank.content_hash,
                "batches": index,
                "human_confirmation_required": True,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return tuple(paths)
